"""
Where customer enquiries get recorded when an operator generates a report.

Two interchangeable backends behind the same two-method interface
(`add_submission` / `update_status`):

- GoogleSheetsSubmissionStore -- the default whenever Google credentials are
  configured (GOOGLE_SHEETS_ID + GOOGLE_SERVICE_ACCOUNT_FILE in Cred.env).
  A hosted app's local disk isn't durable or shared between operators, so
  the enquiry log lives in one online sheet everybody can open via its link.
- ExcelSubmissionStore -- the original local SPG_Customer_Intake_Form.xlsx,
  kept as a fallback so the app still runs on a dev machine with no Google
  setup at all.

`get_submission_store(config)` picks between them; callers never need to
know which one they got.
"""
import datetime
import os
from copy import copy

import openpyxl

from components.config import AppConfig

# Same 25 columns, same order, as the "Customer Enquiries" sheet of
# SPG_Customer_Intake_Form.xlsx -- both backends write identical rows.
HEADERS = [
    "Submission ID", "Date Submitted", "Full Name", "Email Address", "Phone Number",
    "Property Type", "Suburb / Area", "Postcode", "State", "Budget Range", "Buying Purpose",
    "Good Schools Nearby", "Public Transport", "Shopping Centres", "Parks & Green Spaces",
    "Hospital / Medical", "Low Flood Risk", "Low Bushfire Risk", "Quiet Neighbourhood",
    "Investment Potential", "Family Friendly", "Close to CBD",
    "Additional Notes", "Report Status", "Assigned To",
]
STATUS_COL = HEADERS.index("Report Status") + 1  # 1-based, for both openpyxl and gspread

STATUS_GENERATING = "Generating"
STATUS_GENERATED = "Generated"
STATUS_FAILED = "Failed"


def _next_id(last_id) -> str:
    if isinstance(last_id, str) and last_id.startswith("SPG-"):
        try:
            return f"SPG-{int(last_id.split('-')[1]) + 1:03d}"
        except (IndexError, ValueError):
            pass
    return "SPG-001"


def _build_row(submission_id: str, submission: dict, assigned_to: str) -> list:
    return [
        submission_id,
        datetime.date.today().strftime("%d/%m/%Y"),
        submission["full_name"],
        submission["email"],
        submission["phone"],
        submission["property_type"],
        (submission.get("suburb") or "").strip(),
        (submission.get("postcode") or "").strip(),
        (submission.get("state") or "").strip().upper(),
        submission["budget"],
        submission["intention"],
        *submission["priorities_yes_no"],
        submission.get("notes", ""),
        STATUS_GENERATING,
        assigned_to,
    ]


class ExcelSubmissionStore:
    """Local-file fallback. Header lives on row 3, data starts on row 4."""
    HEADER_ROW = 3

    def __init__(self, config: AppConfig):
        self.path = config.excel_path
        self.assigned_to = config.assigned_to
        self.location_label = os.path.basename(self.path)

    def _open(self):
        if not os.path.exists(self.path):
            raise FileNotFoundError(f"Excel file not found: {self.path}")
        wb = openpyxl.load_workbook(self.path)
        return wb, wb.active

    def _last_data_row(self, sheet) -> int:
        last_row = self.HEADER_ROW
        for r in range(self.HEADER_ROW + 1, sheet.max_row + 2):
            if sheet.cell(row=r, column=1).value is not None:
                last_row = r
        return last_row

    def add_submission(self, submission: dict) -> str:
        wb, sheet = self._open()
        last_row = self._last_data_row(sheet)
        next_id = _next_id(sheet.cell(row=last_row, column=1).value)
        next_row = last_row + 1

        for col_idx, val in enumerate(_build_row(next_id, submission, self.assigned_to), start=1):
            new_cell = sheet.cell(row=next_row, column=col_idx, value=val)
            # Carry the previous row's formatting forward so the log stays tidy.
            if last_row > self.HEADER_ROW:
                src_cell = sheet.cell(row=last_row, column=col_idx)
                if src_cell.has_style:
                    new_cell.font = copy(src_cell.font)
                    new_cell.border = copy(src_cell.border)
                    new_cell.fill = copy(src_cell.fill)
                    new_cell.number_format = copy(src_cell.number_format)
                    new_cell.protection = copy(src_cell.protection)
                    new_cell.alignment = copy(src_cell.alignment)

        wb.save(self.path)
        return next_id

    def update_status(self, submission_id: str, status: str) -> None:
        wb, sheet = self._open()
        for r in range(self.HEADER_ROW + 1, sheet.max_row + 1):
            if sheet.cell(row=r, column=1).value == submission_id:
                sheet.cell(row=r, column=STATUS_COL, value=status)
                wb.save(self.path)
                return


class GoogleSheetsSubmissionStore:
    """Online enquiry log: one row per enquiry on the sheet's first tab."""

    def __init__(self, config: AppConfig):
        import gspread
        client = gspread.service_account(filename=config.google_service_account_file)
        self.worksheet = client.open_by_key(config.google_sheets_id).sheet1
        self.assigned_to = config.assigned_to
        self.location_label = "Google Sheets"
        self._ensure_header()

    def _ensure_header(self):
        if self.worksheet.row_values(1) != HEADERS:
            self.worksheet.update([HEADERS], "A1")
            self.worksheet.format("1:1", {"textFormat": {"bold": True}})
            self.worksheet.freeze(rows=1)

    def add_submission(self, submission: dict) -> str:
        ids = self.worksheet.col_values(1)[1:]  # skip header
        next_id = _next_id(ids[-1] if ids else None)
        # RAW keeps values like postcode "0800" or phone "0412..." as text
        # instead of letting Sheets coerce them into numbers.
        self.worksheet.append_row(
            _build_row(next_id, submission, self.assigned_to),
            value_input_option="RAW",
            table_range="A1",
        )
        return next_id

    def update_status(self, submission_id: str, status: str) -> None:
        cell = self.worksheet.find(submission_id, in_column=1)
        if cell:
            self.worksheet.update_cell(cell.row, STATUS_COL, status)


def get_submission_store(config: AppConfig):
    if config.google_sheets_id and config.google_service_account_file:
        return GoogleSheetsSubmissionStore(config)
    return ExcelSubmissionStore(config)
